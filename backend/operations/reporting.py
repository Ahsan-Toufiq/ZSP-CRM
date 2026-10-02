import csv
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape

from django.conf import settings
from django.db.models import IntegerField, Prefetch, Q, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.utils import timezone

from operations.models import AuctionSale, AuctionSaleLine, Container, InventoryBatch


@dataclass(frozen=True)
class InventoryReport:
    generated_at: timezone.datetime
    scope_label: str
    rows: list[dict]
    totals: dict


@dataclass(frozen=True)
class ContainerProfitLossReport:
    generated_at: timezone.datetime
    container: Container
    rows: list[dict]
    totals: dict


def _money(value) -> Decimal:
    rounded = Decimal(value or 0).quantize(Decimal('0.01'))
    return Decimal('0.00') if rounded == 0 else rounded


def _logo_path() -> Path:
    return Path(settings.BASE_DIR) / 'static' / 'branding' / 'digi7-logo.png'


def _content_disposition(filename: str, *, inline: bool = False) -> str:
    disposition = 'inline' if inline else 'attachment'
    return f'{disposition}; filename="{filename}"'


def _report_header(story, styles, title: str, subtitle: str, generated_at, *, generated_label='Prepared on', compact=False):
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, Spacer, Table, TableStyle

    logo_path = _logo_path()
    header_data = []
    if logo_path.exists():
        header_data.append(Image(str(logo_path), width=0.62 * inch, height=0.62 * inch))
    else:
        header_data.append(Paragraph('<b>ZSP</b>', styles['Title']))
    header_data.append(Paragraph(f'<b>{title}</b><br/><font size="9">{subtitle}</font>', styles['Title']))
    header_data.append(Paragraph(f'{generated_label}<br/><b>{generated_at.strftime("%Y-%m-%d %H:%M:%S %Z")}</b>', styles['Normal']))
    widths = [0.75 * inch, 4.15 * inch, 2.25 * inch] if compact else [0.8 * inch, 6.7 * inch, 3.3 * inch]
    header = Table([header_data], colWidths=widths)
    header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f3faf7')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.HexColor('#0f766e')),
        ('INNERPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.extend([header, Spacer(1, 12)])


def build_inventory_report(*, container_id=None) -> InventoryReport:
    generated_at = timezone.localtime()
    queryset = (
        InventoryBatch.objects
        .select_related('item', 'container', 'container_item')
        .annotate(
            sold_quantity_total=Coalesce(
                Sum('sale_lines__quantity', filter=Q(sale_lines__sale__is_cancelled=False)),
                Value(0),
                output_field=IntegerField(),
            ),
        )
        .order_by('container__reference', 'item__part_name', 'item__part_number', 'item__category', 'created_at')
    )
    if container_id:
        queryset = queryset.filter(container_id=container_id)

    rows = []
    for batch in queryset:
        available = max(int(batch.quantity) - int(batch.sold_quantity_total or 0), 0)
        if available <= 0:
            continue
        raw_total = _money(batch.raw_unit_cost) * Decimal(available)
        net_total = _money(batch.net_unit_cost) * Decimal(available)
        rows.append({
            'part_name': batch.item.part_name,
            'part_number': batch.item.part_number or '',
            'category': batch.item.category or '',
            'unit': batch.item.unit or 'piece',
            'container_reference': batch.container.reference if batch.container_id else batch.source_label or 'General inventory',
            'container_status': batch.container.get_status_display() if batch.container_id else '',
            'source_label': batch.source_label or '',
            'batch_quantity': batch.quantity,
            'sold_quantity': batch.sold_quantity_total or 0,
            'available_quantity': available,
            'raw_unit_cost': _money(batch.raw_unit_cost),
            'net_unit_cost': _money(batch.net_unit_cost),
            'added_cost_share_unit': _money(batch.net_unit_cost) - _money(batch.raw_unit_cost),
            'raw_total': raw_total,
            'net_total': net_total,
            'added_cost_share_total': net_total - raw_total,
            'notes': batch.notes or '',
        })

    scope_label = 'All available inventory'
    if container_id and rows:
        scope_label = f"Available inventory for {rows[0]['container_reference']}"
    elif container_id:
        scope_label = 'Available inventory for selected container'

    return InventoryReport(
        generated_at=generated_at,
        scope_label=scope_label,
        rows=rows,
        totals={
            'row_count': len(rows),
            'available_quantity': sum(row['available_quantity'] for row in rows),
            'raw_total': sum((row['raw_total'] for row in rows), Decimal('0.00')),
            'net_total': sum((row['net_total'] for row in rows), Decimal('0.00')),
            'added_cost_share_total': sum((row['added_cost_share_total'] for row in rows), Decimal('0.00')),
        },
    )


def inventory_report_csv_response(report: InventoryReport) -> HttpResponse:
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(['ZSP Available Inventory Report'])
    writer.writerow(['Scope', report.scope_label])
    writer.writerow(['Generated at', report.generated_at.strftime('%Y-%m-%d %H:%M:%S %Z')])
    writer.writerow([])
    writer.writerow(['Summary'])
    writer.writerow(['Rows', report.totals['row_count']])
    writer.writerow(['Available quantity', report.totals['available_quantity']])
    writer.writerow(['Raw total', report.totals['raw_total']])
    writer.writerow(['Clearing fee allocation', report.totals['added_cost_share_total']])
    writer.writerow(['Net total', report.totals['net_total']])
    writer.writerow([])
    writer.writerow([
        'Part name', 'Part number', 'Category', 'Unit', 'Source container', 'Container status',
        'Batch quantity', 'Sold quantity', 'Available quantity', 'Raw unit cost', 'Clearing fee / unit',
        'Net unit cost', 'Raw total', 'Clearing fee allocation', 'Net total', 'Notes',
    ])
    for row in report.rows:
        writer.writerow([
            row['part_name'], row['part_number'], row['category'], row['unit'],
            row['container_reference'], row['container_status'], row['batch_quantity'],
            row['sold_quantity'], row['available_quantity'], row['raw_unit_cost'],
            row['added_cost_share_unit'], row['net_unit_cost'], row['raw_total'],
            row['added_cost_share_total'], row['net_total'], row['notes'],
        ])

    response = HttpResponse(buffer.getvalue(), content_type='text/csv')
    response['Content-Disposition'] = _content_disposition('zsp-available-inventory.csv')
    return response


def inventory_report_pdf_response(report: InventoryReport) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    story = []
    _report_header(story, styles, 'ZSP Available Inventory Report', report.scope_label, report.generated_at)

    summary = Table([[
        'Rows', report.totals['row_count'],
        'Available Qty', report.totals['available_quantity'],
        'Raw Total', f"PKR {report.totals['raw_total']:,.2f}",
        'Net Total', f"PKR {report.totals['net_total']:,.2f}",
    ]])
    summary.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#0f766e')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('INNERPADDING', (0, 0), (-1, -1), 7),
    ]))
    story.extend([summary, Spacer(1, 10)])

    rows = [[
        'Part', 'No.', 'Category', 'Source', 'Avail', 'Raw Unit', 'Fee/Unit',
        'Net Unit', 'Raw Total', 'Fee Share', 'Net Total',
    ]]
    for row in report.rows:
        rows.append([
            Paragraph(escape(row['part_name']), styles['BodyText']),
            Paragraph(escape(row['part_number'] or '-'), styles['BodyText']),
            Paragraph(escape(row['category'] or '-'), styles['BodyText']),
            Paragraph(escape(row['container_reference']), styles['BodyText']),
            f"{row['available_quantity']} {row['unit']}",
            f"PKR {row['raw_unit_cost']:,.2f}",
            f"PKR {row['added_cost_share_unit']:,.2f}",
            f"PKR {row['net_unit_cost']:,.2f}",
            f"PKR {row['raw_total']:,.2f}",
            f"PKR {row['added_cost_share_total']:,.2f}",
            f"PKR {row['net_total']:,.2f}",
        ])
    if len(rows) == 1:
        rows.append(['No available inventory', '', '', '', '', '', '', '', '', '', ''])

    table = Table(
        rows,
        repeatRows=1,
        colWidths=[1.55 * inch, 0.85 * inch, 0.95 * inch, 1.05 * inch, 0.58 * inch, 0.82 * inch, 0.82 * inch, 0.82 * inch, 0.9 * inch, 0.9 * inch, 0.95 * inch],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#111827')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 6.5),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    story.append(table)
    doc.build(story)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition('zsp-available-inventory.pdf')
    return response


def build_container_profit_loss_report(*, container: Container) -> ContainerProfitLossReport:
    sale_lines = (
        AuctionSaleLine.objects
        .filter(sale__is_cancelled=False)
        .select_related('sale', 'sale__customer')
        .order_by('sale__sale_date', 'sale__created_at')
    )
    batches = (
        InventoryBatch.objects
        .filter(container=container)
        .select_related('item', 'container_item')
        .prefetch_related(Prefetch('sale_lines', queryset=sale_lines))
        .order_by('item__part_name', 'item__part_number', 'created_at')
    )

    rows = []
    for batch in batches:
        lines = list(batch.sale_lines.all())
        sold_quantity = sum(int(line.quantity) for line in lines)
        available_quantity = max(int(batch.quantity) - sold_quantity, 0)
        revenue = sum((Decimal(line.sold_price) * Decimal(line.quantity) for line in lines), Decimal('0.00'))
        sold_cost = _money(batch.net_unit_cost) * Decimal(sold_quantity)
        remaining_value = _money(batch.net_unit_cost) * Decimal(available_quantity)
        rows.append({
            'part_name': batch.item.part_name,
            'part_number': batch.item.part_number or '',
            'category': batch.item.category or '',
            'unit': batch.item.unit or 'piece',
            'acquired_quantity': int(batch.quantity),
            'sold_quantity': sold_quantity,
            'available_quantity': available_quantity,
            'raw_unit_cost': _money(batch.raw_unit_cost),
            'net_unit_cost': _money(batch.net_unit_cost),
            'revenue': _money(revenue),
            'sold_cost': _money(sold_cost),
            'gross_profit': _money(revenue - sold_cost),
            'remaining_value': _money(remaining_value),
        })

    raw_inventory_cost = sum(
        (row['raw_unit_cost'] * Decimal(row['acquired_quantity']) for row in rows),
        Decimal('0.00'),
    )
    allocated_inventory_cost = sum(
        (row['net_unit_cost'] * Decimal(row['acquired_quantity']) for row in rows),
        Decimal('0.00'),
    )
    revenue = sum((row['revenue'] for row in rows), Decimal('0.00'))
    sold_cost = sum((row['sold_cost'] for row in rows), Decimal('0.00'))
    remaining_value = sum((row['remaining_value'] for row in rows), Decimal('0.00'))
    total_recorded_cost = raw_inventory_cost + _money(container.added_cost)

    return ContainerProfitLossReport(
        generated_at=timezone.localtime(),
        container=container,
        rows=rows,
        totals={
            'raw_inventory_cost': _money(raw_inventory_cost),
            'clearing_fee': _money(container.added_cost),
            'total_recorded_cost': _money(total_recorded_cost),
            'allocated_inventory_cost': _money(allocated_inventory_cost),
            'allocation_variance': _money(total_recorded_cost - allocated_inventory_cost),
            'revenue': _money(revenue),
            'sold_cost': _money(sold_cost),
            'gross_profit': _money(revenue - sold_cost),
            'gross_margin': (revenue - sold_cost) / revenue * Decimal('100.00') if revenue else Decimal('0.00'),
            'remaining_inventory_value': _money(remaining_value),
            'sold_quantity': sum(row['sold_quantity'] for row in rows),
            'available_quantity': sum(row['available_quantity'] for row in rows),
        },
    )


def container_profit_loss_pdf_response(report: ContainerProfitLossReport) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), rightMargin=22, leftMargin=22, topMargin=20, bottomMargin=22,
    )
    styles = getSampleStyleSheet()
    cell = ParagraphStyle('ContainerPnlCell', parent=styles['BodyText'], fontSize=6.8, leading=8.2, wordWrap='CJK')
    note = ParagraphStyle('ContainerPnlNote', parent=styles['BodyText'], fontSize=7.5, leading=10, textColor=colors.HexColor('#475569'))
    story = []
    container = report.container
    _report_header(
        story,
        styles,
        'Container Profit & Loss',
        f'Syed Zulfiqar Old Spare Parts | {escape(container.reference)}',
        report.generated_at,
    )

    details = Table([
        ['Container', container.reference, 'Status', container.get_status_display(), 'Current location', container.current_location or '-'],
        ['Agent', container.supplier_name or '-', 'Arrival date', str(container.arrival_date or '-'), 'Size / Type', container.size_type or '-'],
    ], colWidths=[0.9 * inch, 2.15 * inch, 0.8 * inch, 1.45 * inch, 1.15 * inch, 3.05 * inch])
    details.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3faf7')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#f3faf7')),
        ('BACKGROUND', (4, 0), (4, -1), colors.HexColor('#f3faf7')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTNAME', (4, 0), (4, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERPADDING', (0, 0), (-1, -1), 6),
    ]))

    totals = report.totals
    summary = Table([
        ['Recorded container cost', f"PKR {totals['total_recorded_cost']:,.0f}", 'Sales revenue', f"PKR {totals['revenue']:,.0f}", 'Realized COGS', f"PKR {totals['sold_cost']:,.0f}"],
        ['Realized gross profit', f"PKR {totals['gross_profit']:,.0f}", 'Gross margin', f"{totals['gross_margin']:,.1f}%", 'Remaining stock value', f"PKR {totals['remaining_inventory_value']:,.0f}"],
        ['Raw parts cost', f"PKR {totals['raw_inventory_cost']:,.0f}", 'Clearing Fee', f"PKR {totals['clearing_fee']:,.0f}", 'Allocation variance', f"PKR {totals['allocation_variance']:,.0f}"],
    ], colWidths=[1.35 * inch, 1.65 * inch, 1.15 * inch, 1.55 * inch, 1.35 * inch, 1.65 * inch])
    summary.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f766e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('INNERPADDING', (0, 0), (-1, -1), 7),
    ]))
    story.extend([details, Spacer(1, 10), summary, Spacer(1, 10)])

    rows = [['Part', 'Part no.', 'Category', 'Acquired', 'Sold', 'Available', 'Net unit', 'Revenue', 'COGS', 'Gross P/L', 'Stock value']]
    for row in report.rows:
        rows.append([
            Paragraph(escape(row['part_name']), cell),
            Paragraph(escape(row['part_number'] or '-'), cell),
            Paragraph(escape(row['category'] or '-'), cell),
            f"{row['acquired_quantity']} {row['unit']}",
            str(row['sold_quantity']),
            str(row['available_quantity']),
            f"PKR {row['net_unit_cost']:,.0f}",
            f"PKR {row['revenue']:,.0f}",
            f"PKR {row['sold_cost']:,.0f}",
            f"PKR {row['gross_profit']:,.0f}",
            f"PKR {row['remaining_value']:,.0f}",
        ])
    if len(rows) == 1:
        rows.append(['No inventory has been recorded for this container.', '', '', '', '', '', '', '', '', '', ''])
    table = Table(
        rows,
        repeatRows=1,
        colWidths=[1.45 * inch, 0.8 * inch, 0.9 * inch, 0.65 * inch, 0.48 * inch, 0.58 * inch, 0.82 * inch, 0.85 * inch, 0.8 * inch, 0.8 * inch, 0.85 * inch],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#111827')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 6.5),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('INNERPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.extend([
        table,
        Spacer(1, 10),
        KeepTogether(Paragraph(
            'This is a gross container-level report. Cost of goods sold uses the current allocated net cost of each source batch. General operating expenses are not allocated to containers and are therefore excluded.',
            note,
        )),
    ])
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition(f'container-{container.reference}-profit-loss.pdf')
    return response


def _auction_inventory_units(report: InventoryReport) -> list[dict]:
    units = []
    serial = 1
    for row in report.rows:
        quantity = int(row['available_quantity'])
        for unit_number in range(1, quantity + 1):
            units.append({
                'serial': serial,
                'part_name': row['part_name'],
                'category': row['category'],
                'unit': row['unit'],
                'unit_number': unit_number,
                'quantity': quantity,
            })
            serial += 1
    return units


def auction_inventory_sheet_pdf_response(container: Container) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    report = build_inventory_report(container_id=container.id)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=28, leftMargin=28, topMargin=24, bottomMargin=28)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('AuctionInventoryCell', parent=styles['BodyText'], fontSize=8.5, leading=11)
    story = []
    _report_header(
        story, styles, 'ZSP Auction Inventory Sheet',
        f'Container {escape(container.reference)}', report.generated_at, compact=True,
    )
    details = Table([
        ['Container Number', Paragraph(escape(container.reference), cell_style), 'Size / Type', Paragraph(escape(container.size_type or '-'), cell_style)],
        ['Current Location', Paragraph(escape(container.current_location or '-'), cell_style), 'Agent', Paragraph(escape(container.supplier_name or '-'), cell_style)],
    ], colWidths=[1.2 * inch, 2.35 * inch, 1.15 * inch, 2.4 * inch])
    details.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#d1d5db')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3faf7')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#f3faf7')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERPADDING', (0, 0), (-1, -1), 7),
    ]))
    story.extend([details, Spacer(1, 12)])

    rows = [['S.No.', 'Item / Part Name', 'Category', 'Unit', 'Buyer Name', 'Selling Price']]
    for unit in _auction_inventory_units(report):
        rows.append([
            str(unit['serial']),
            Paragraph(escape(unit['part_name']), cell_style),
            Paragraph(escape(unit['category'] or '-'), cell_style),
            f"{unit['unit_number']} of {unit['quantity']} ({unit['unit']})",
            '',
            '',
        ])
    if len(rows) == 1:
        rows.append(['', 'No available inventory', '', '', '', ''])
    table = Table(
        rows,
        repeatRows=1,
        colWidths=[0.42 * inch, 1.82 * inch, 1.08 * inch, 0.95 * inch, 1.45 * inch, 1.35 * inch],
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#111827')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.45, colors.HexColor('#9ca3af')),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 11),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 11),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    totals = Table([
        [Paragraph('<b>Auction Totals</b>', styles['Heading3']), '', '', ''],
        ['Total Auction Sales', '', 'Total Inventory Cost', ''],
    ], colWidths=[1.55 * inch, 2.0 * inch, 1.35 * inch, 2.17 * inch], rowHeights=[0.38 * inch, 0.72 * inch])
    totals.setStyle(TableStyle([
        ('SPAN', (0, 0), (-1, 0)),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f3faf7')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0f4f4b')),
        ('BOX', (0, 0), (-1, -1), 0.7, colors.HexColor('#64748b')),
        ('GRID', (0, 1), (-1, -1), 0.45, colors.HexColor('#94a3b8')),
        ('FONTNAME', (0, 1), (0, 1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 1), (2, 1), 'Helvetica-Bold'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.extend([table, Spacer(1, 14), totals])
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition(f'zsp-auction-inventory-{container.reference}.pdf')
    return response


def _tracking_rows(containers):
    return [
        {
            'serial': index,
            'reference': container.reference,
            'size_type': container.size_type or '',
            'current_location': container.current_location or '',
            'agent': container.supplier_name or '',
            'notes': container.notes or '',
        }
        for index, container in enumerate(containers, start=1)
    ]


def container_tracking_pdf_response(containers) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    generated_at = timezone.localtime()
    rows_data = _tracking_rows(containers)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=22, bottomMargin=24)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('TrackingCell', parent=styles['BodyText'], fontSize=7.5, leading=9.5)
    story = []
    _report_header(story, styles, 'ZSP Container Tracking Report', 'Container movement and status register', generated_at)
    rows = [['Serial No.', 'Container Number', 'Size / Type', 'Current Location', 'Agent', 'Notes']]
    for row in rows_data:
        rows.append([
            row['serial'],
            Paragraph(escape(row['reference']), cell_style),
            Paragraph(escape(row['size_type'] or '-'), cell_style),
            Paragraph(escape(row['current_location'] or '-'), cell_style),
            Paragraph(escape(row['agent'] or '-'), cell_style),
            Paragraph(escape(row['notes'] or '-'), cell_style),
        ])
    if len(rows) == 1:
        rows.append(['No containers recorded', '', '', '', '', ''])
    table = Table(rows, repeatRows=1, colWidths=[0.62 * inch, 1.55 * inch, 1.25 * inch, 1.75 * inch, 1.75 * inch, 3.65 * inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#111827')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#d1d5db')),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERPADDING', (0, 0), (-1, -1), 6),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    story.append(table)
    doc.build(story)
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition('zsp-container-tracking.pdf')
    return response


def container_tracking_xlsx_response(containers) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorkbookImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rows = _tracking_rows(containers)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Container Tracking'
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells('B1:F1')
    sheet['B1'] = 'ZSP Container Tracking Report'
    sheet['B1'].font = Font(size=18, bold=True, color='111827')
    sheet['B1'].alignment = Alignment(vertical='center')
    sheet.merge_cells('B2:F2')
    sheet['B2'] = f"Prepared on {timezone.localtime().strftime('%d %B %Y, %I:%M %p')}"
    sheet['B2'].font = Font(size=10, color='4B5563')
    logo_path = _logo_path()
    if logo_path.exists():
        logo = WorkbookImage(str(logo_path))
        logo.width = 64
        logo.height = 64
        sheet.add_image(logo, 'A1')
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 22
    headers = ['Serial No.', 'Container Number', 'Size / Type', 'Current Location', 'Agent', 'Notes']
    header_row = 4
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='111827')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='D1D5DB')
    for row_index, row in enumerate(rows, start=header_row + 1):
        values = [row['serial'], row['reference'], row['size_type'], row['current_location'], row['agent'], row['notes']]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if row_index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F8FAFC')
        sheet.row_dimensions[row_index].height = 34
    for cell in sheet[header_row]:
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = {'A': 12, 'B': 24, 'C': 18, 'D': 28, 'E': 26, 'F': 58}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A5'
    sheet.auto_filter.ref = f'A4:F{max(header_row, header_row + len(rows))}'
    sheet.print_title_rows = '1:4'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = _content_disposition('zsp-container-tracking.xlsx')
    return response


def sale_invoice_pdf_response(sale: AuctionSale, *, inline: bool = False) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    sale = (
        AuctionSale.objects
        .select_related('customer', 'gate_pass')
        .prefetch_related('lines__item', 'lines__inventory_batch__container', 'cheques')
        .get(id=sale.id)
    )
    generated_at = timezone.localtime()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=26, bottomMargin=26)
    styles = getSampleStyleSheet()
    meta_style = ParagraphStyle('InvoiceMetaCell', parent=styles['BodyText'], fontSize=8, leading=10)
    story = []
    _report_header(
        story, styles, 'Syed Zulfiqar Old Spare Parts',
        f'Sales Invoice {sale.sale_number}', generated_at, generated_label='Date / time', compact=True,
    )

    customer_name = sale.customer.name if sale.customer_id else 'Cash customer'
    customer_phone = sale.customer.phone if sale.customer_id else '-'
    try:
        gate_pass_number = sale.gate_pass.gate_pass_number
    except AuctionSale.gate_pass.RelatedObjectDoesNotExist:
        gate_pass_number = '-'
    invoice_meta = Table([
        ['Invoice #', Paragraph(escape(sale.sale_number), meta_style), 'Date', str(sale.sale_date)],
        ['Customer', Paragraph(escape(customer_name), meta_style), 'Contact', Paragraph(escape(customer_phone), meta_style)],
        ['Payment', Paragraph(escape(sale.get_payment_type_display()), meta_style), 'Gate pass', Paragraph(escape(gate_pass_number), meta_style)],
    ], colWidths=[1.2 * inch, 2.4 * inch, 1.2 * inch, 2.1 * inch])
    invoice_meta.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3faf7')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#f3faf7')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.extend([invoice_meta, Spacer(1, 14)])

    item_rows = [['Item', 'Part No.', 'Category', 'Quantity', 'Unit Price', 'Amount']]
    for line in sale.lines.all():
        item_rows.append([
            Paragraph(escape(line.item.part_name), styles['BodyText']),
            Paragraph(escape(line.item.part_number or '-'), styles['BodyText']),
            Paragraph(escape(line.item.category or '-'), styles['BodyText']),
            f'{line.quantity} {line.item.unit}',
            f'PKR {line.sold_price:,.2f}',
            f'PKR {(line.quantity * line.sold_price):,.2f}',
        ])
    items = Table(item_rows, repeatRows=1, colWidths=[2.0 * inch, 1.0 * inch, 1.05 * inch, 0.9 * inch, 1.0 * inch, 1.1 * inch])
    items.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#111827')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.6),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    story.extend([items, Spacer(1, 12)])

    cheque_total = sum((cheque.amount for cheque in sale.cheques.all()), Decimal('0.00'))
    totals = Table([
        ['Items total', f'PKR {sale.total_amount:,.2f}'],
        ['Cash paid', f'PKR {sale.cash_amount:,.2f}'],
        ['Cheque paid', f'PKR {cheque_total:,.2f}'],
        ['Credit amount', f'PKR {sale.credit_amount:,.2f}'],
        ['Balance due', f'PKR {sale.receivable_amount:,.2f}'],
        ['Invoice amount', f'PKR {sale.total_amount:,.2f}'],
    ], colWidths=[2.0 * inch, 1.4 * inch], hAlign='RIGHT')
    totals.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#d1d5db')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3faf7')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#0f766e')),
        ('TEXTCOLOR', (0, -1), (-1, -1), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('INNERPADDING', (0, 0), (-1, -1), 7),
    ]))
    story.extend([totals, Spacer(1, 18)])
    story.append(Paragraph('Thank you for your business. This invoice was issued by Syed Zulfiqar Old Spare Parts through Digi7.', styles['BodyText']))
    doc.build(story)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition(f'zsp-invoice-{sale.sale_number}.pdf', inline=inline)
    return response


def sale_thermal_invoice_pdf_response(sale: AuctionSale, *, inline: bool = False) -> HttpResponse:
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    sale = (
        AuctionSale.objects
        .select_related('customer')
        .prefetch_related('lines__item', 'cheques')
        .get(id=sale.id)
    )
    page_height = max(150 * mm, (118 + (len(sale.lines.all()) * 11)) * mm)
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=(80 * mm, page_height),
        rightMargin=5 * mm,
        leftMargin=5 * mm,
        topMargin=5 * mm,
        bottomMargin=5 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('ThermalTitle', parent=styles['Title'], fontSize=13, leading=15, alignment=1, spaceAfter=2)
    center_style = ParagraphStyle('ThermalCenter', parent=styles['BodyText'], fontSize=7.5, leading=9, alignment=1)
    cell_style = ParagraphStyle('ThermalCell', parent=styles['BodyText'], fontSize=7, leading=8.5)
    story = []
    logo_path = _logo_path()
    if logo_path.exists():
        logo = Image(str(logo_path), width=16 * mm, height=16 * mm)
        logo.hAlign = 'CENTER'
        story.append(logo)
    story.extend([
        Paragraph('<b>SYED ZULFIQAR OLD SPARE PARTS</b>', title_style),
        Paragraph('Spare Parts Auction', center_style),
        Spacer(1, 3 * mm),
    ])

    customer_name = sale.customer.name if sale.customer_id else 'Cash customer'
    customer_phone = sale.customer.phone if sale.customer_id else '-'
    meta = Table([
        ['Invoice', escape(sale.sale_number)],
        ['Date', sale.sale_date.strftime('%d %b %Y')],
        ['Customer', Paragraph(escape(customer_name), cell_style)],
        ['Contact', escape(customer_phone)],
        ['Payment', sale.get_payment_type_display()],
    ], colWidths=[18 * mm, 52 * mm])
    meta.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
    ]))
    story.extend([meta, Spacer(1, 3 * mm)])

    rows = [['Item', 'Qty', 'Rate', 'Amount']]
    for line in sale.lines.all():
        rows.append([
            Paragraph(
                f"<b>{escape(line.item.part_name)}</b>"
                f"<br/><font size='6'>{escape(line.item.part_number or line.item.category or '')}</font>",
                cell_style,
            ),
            f'{line.quantity}',
            f'{line.sold_price:,.0f}',
            f'{(line.quantity * line.sold_price):,.0f}',
        ])
    items = Table(rows, repeatRows=1, colWidths=[34 * mm, 9 * mm, 13 * mm, 14 * mm])
    items.setStyle(TableStyle([
        ('LINEABOVE', (0, 0), (-1, 0), 0.7, colors.black),
        ('LINEBELOW', (0, 0), (-1, 0), 0.7, colors.black),
        ('LINEBELOW', (0, -1), (-1, -1), 0.7, colors.black),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 6.7),
        ('ALIGN', (1, 1), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 1.5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 1.5),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.extend([items, Spacer(1, 3 * mm)])

    cheque_total = sum((cheque.amount for cheque in sale.cheques.all()), Decimal('0.00'))
    totals = Table([
        ['Invoice total', f'PKR {sale.total_amount:,.0f}'],
        ['Cash paid', f'PKR {sale.cash_amount:,.0f}'],
        ['Cheque', f'PKR {cheque_total:,.0f}'],
        ['Balance due', f'PKR {sale.receivable_amount:,.0f}'],
    ], colWidths=[35 * mm, 35 * mm])
    totals.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('LINEABOVE', (0, 0), (-1, 0), 0.7, colors.black),
        ('LINEBELOW', (0, -1), (-1, -1), 0.7, colors.black),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.extend([
        totals,
        Spacer(1, 4 * mm),
        Paragraph('Thank you for your business.', center_style),
        Paragraph('Syed Zulfiqar Old Spare Parts powered by Digi7', center_style),
    ])
    doc.build(story)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = _content_disposition(f'zsp-thermal-invoice-{sale.sale_number}.pdf', inline=inline)
    return response
