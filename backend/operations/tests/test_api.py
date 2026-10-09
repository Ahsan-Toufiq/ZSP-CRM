from decimal import Decimal
from datetime import timedelta
from io import BytesIO

import pytest
from django.contrib.auth.models import User
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.models import UserProfile
from accounts.permissions import full_tab_permissions
from finance.models import CustomerLedgerEntry
from operations.models import AuctionSale, Container, ContainerItem, ContainerStatusAppearance, Customer, InventoryBatch, PartInventory
from operations.reporting import _auction_inventory_units, build_container_profit_loss_report, build_inventory_report
from operations.services import create_auction_sale


@pytest.fixture
def admin_user(db):
    user = User.objects.create_user(username='api-admin', password='StrongPass123!')
    UserProfile.objects.create(user=user, tab_permissions=full_tab_permissions())
    return user


@pytest.fixture
def api_client(admin_user):
    client = APIClient()
    client.login(username=admin_user.username, password='StrongPass123!')
    return client


@pytest.mark.django_db
def test_customer_without_transactions_can_be_deleted(api_client):
    customer = Customer.objects.create(name='Delete Me', phone='+923001231231')

    response = api_client.delete(f'/api/operations/customers/{customer.id}/')

    assert response.status_code == 204
    assert not Customer.objects.filter(id=customer.id).exists()


@pytest.mark.django_db
def test_container_status_colors_return_defaults_and_persist_overrides(api_client):
    defaults = api_client.get('/api/operations/container-status-appearances/')

    assert defaults.status_code == 200
    assert len(defaults.data) == len(Container.Status.choices)
    ready = next(item for item in defaults.data if item['status'] == Container.Status.READY_FOR_AUCTION)
    assert ready == {
        'status': Container.Status.READY_FOR_AUCTION,
        'label': 'Ready for auction',
        'color': '#16A34A',
        'is_custom': False,
    }

    updated = api_client.post(
        '/api/operations/container-status-appearances/',
        {'status': Container.Status.READY_FOR_AUCTION, 'color': '#123abc'},
        format='json',
    )

    assert updated.status_code == 200
    assert updated.data['color'] == '#123ABC'
    assert updated.data['is_custom'] is True
    assert ContainerStatusAppearance.objects.get(status=Container.Status.READY_FOR_AUCTION).color == '#123ABC'


@pytest.mark.django_db
def test_invalid_container_status_color_does_not_replace_existing_override(api_client):
    ContainerStatusAppearance.objects.create(status=Container.Status.CLOSED, color='#112233')

    response = api_client.post(
        '/api/operations/container-status-appearances/',
        {'status': Container.Status.CLOSED, 'color': 'not-a-color'},
        format='json',
    )

    assert response.status_code == 400
    assert ContainerStatusAppearance.objects.get(status=Container.Status.CLOSED).color == '#112233'


@pytest.mark.django_db
def test_customer_type_defaults_when_omitted(api_client):
    response = api_client.post(
        '/api/operations/customers/',
        {'name': 'Default Type Customer', 'phone': '+923001231233'},
        format='json',
    )

    assert response.status_code == 201
    assert response.data['customer_type'] == Customer.CustomerType.INDIVIDUAL


@pytest.mark.django_db
def test_customer_opening_balance_creates_receivable_ledger_entry(api_client):
    response = api_client.post(
        '/api/operations/customers/',
        {
            'name': 'Opening Balance Customer',
            'phone': '+923001231234',
            'opening_balance': '25000.00',
            'opening_balance_direction': 'receivable',
        },
        format='json',
    )

    assert response.status_code == 201
    customer = Customer.objects.get(id=response.data['id'])
    entry = CustomerLedgerEntry.objects.get(customer=customer)
    assert entry.entry_type == CustomerLedgerEntry.EntryType.ADJUSTMENT
    assert entry.description == 'Opening balance'
    assert entry.debit == Decimal('25000.00')
    assert entry.credit == Decimal('0.00')


@pytest.mark.django_db
def test_customer_opening_balance_can_record_customer_credit(api_client):
    response = api_client.post(
        '/api/operations/customers/',
        {
            'name': 'Advance Customer',
            'phone': '+923001231235',
            'opening_balance': '5000.00',
            'opening_balance_direction': 'credit',
        },
        format='json',
    )

    assert response.status_code == 201
    entry = CustomerLedgerEntry.objects.get(customer_id=response.data['id'])
    assert entry.debit == Decimal('0.00')
    assert entry.credit == Decimal('5000.00')


@pytest.mark.django_db
def test_customer_without_opening_balance_has_no_ledger_entry(api_client):
    response = api_client.post(
        '/api/operations/customers/',
        {'name': 'No Opening Balance', 'phone': '+923001231236'},
        format='json',
    )

    assert response.status_code == 201
    assert CustomerLedgerEntry.objects.filter(customer_id=response.data['id']).count() == 0


@pytest.mark.django_db
def test_customer_opening_balance_can_be_edited(api_client):
    response = api_client.post(
        '/api/operations/customers/',
        {
            'name': 'Editable Opening Balance',
            'phone': '+923001231237',
            'opening_balance': '25000.00',
            'opening_balance_direction': 'receivable',
        },
        format='json',
    )
    assert response.status_code == 201

    update = api_client.patch(
        f"/api/operations/customers/{response.data['id']}/",
        {
            'opening_balance': '8000.00',
            'opening_balance_direction': 'credit',
        },
        format='json',
    )

    assert update.status_code == 200
    entry = CustomerLedgerEntry.objects.get(customer_id=response.data['id'])
    assert entry.debit == Decimal('0.00')
    assert entry.credit == Decimal('8000.00')
    assert update.data['opening_balance'] == '8000.00'
    assert update.data['opening_balance_direction'] == 'credit'


@pytest.mark.django_db
def test_customer_opening_balance_can_be_removed(api_client):
    response = api_client.post(
        '/api/operations/customers/',
        {
            'name': 'Removable Opening Balance',
            'phone': '+923001231238',
            'opening_balance': '12000.00',
        },
        format='json',
    )

    update = api_client.patch(
        f"/api/operations/customers/{response.data['id']}/",
        {'opening_balance': '0.00'},
        format='json',
    )

    assert update.status_code == 200
    assert CustomerLedgerEntry.objects.filter(customer_id=response.data['id']).count() == 0


@pytest.mark.django_db
def test_customer_with_transactions_cannot_be_deleted(api_client, admin_user):
    customer = Customer.objects.create(name='Keep Me', phone='+923001231232')
    item = PartInventory.objects.create(part_name='Door', quantity=1, unit='piece')
    batch = InventoryBatch.objects.create(item=item, quantity=1, raw_unit_cost=Decimal('5000.00'), source_label='Test batch')
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CREDIT,
        customer=customer,
        lines=[{'inventory_batch': batch, 'quantity': 1, 'sold_price': Decimal('15000.00')}],
    )

    response = api_client.delete(f'/api/operations/customers/{customer.id}/')

    assert response.status_code == 409
    assert response.data['blocking_records']['auction_sales'] == 1
    assert Customer.objects.filter(id=customer.id).exists()


@pytest.mark.django_db
def test_empty_container_can_be_deleted(api_client):
    container = Container.objects.create(reference='DELETE-EMPTY')

    response = api_client.delete(f'/api/operations/containers/{container.id}/')

    assert response.status_code == 204
    assert not Container.objects.filter(id=container.id).exists()


@pytest.mark.django_db
def test_container_with_inventory_returns_conflict_on_delete(api_client):
    container = Container.objects.create(reference='KEEP-WITH-STOCK')
    ContainerItem.objects.create(
        container=container,
        part_name='Mirror',
        category='Body',
        quantity=1,
        unit='piece',
        raw_unit_cost=Decimal('2000.00'),
    )

    response = api_client.delete(f'/api/operations/containers/{container.id}/')

    assert response.status_code == 409
    assert response.data['protected_records'] == 1
    assert Container.objects.filter(id=container.id).exists()


@pytest.mark.django_db
def test_part_can_be_created_with_general_inventory_only(api_client):
    response = api_client.post(
        '/api/operations/parts/',
        {
            'part_name': 'General stock alternator',
            'part_number': 'GEN-ALT-1',
            'category': 'Electrical',
            'unit': 'piece',
            'sources': [{
                'source_type': 'general',
                'source_label': 'Local market purchase',
                'quantity': 5,
                'raw_unit_cost': '12000.00',
                'description': 'Five units bought together.',
            }],
        },
        format='json',
    )

    assert response.status_code == 201
    item = PartInventory.objects.get(id=response.data['id'])
    batch = item.batches.get()
    assert item.quantity == 5
    assert batch.container_id is None
    assert batch.container_item_id is None
    assert batch.quantity == 5
    assert batch.raw_unit_cost == Decimal('12000.00')
    assert batch.net_unit_cost == Decimal('12000.00')


@pytest.mark.django_db
def test_general_inventory_merges_matching_batch_and_separates_different_cost(api_client):
    item = PartInventory.objects.create(
        part_name='General mirror', part_number='GM-1', category='Body', quantity=0, unit='piece',
    )
    base_payload = {
        'item': str(item.id),
        'source_label': 'Local supplier',
        'quantity': 5,
        'raw_unit_cost': '2000.00',
        'notes': 'Same acquisition details',
    }
    first = api_client.post('/api/operations/inventory-batches/', base_payload, format='json')
    second = api_client.post(
        '/api/operations/inventory-batches/',
        {**base_payload, 'quantity': 2},
        format='json',
    )
    different_cost = api_client.post(
        '/api/operations/inventory-batches/',
        {**base_payload, 'quantity': 3, 'raw_unit_cost': '2600.00'},
        format='json',
    )

    assert first.status_code == second.status_code == different_cost.status_code == 201
    item.refresh_from_db()
    assert item.quantity == 10
    assert item.batches.count() == 2
    assert item.batches.get(raw_unit_cost=Decimal('2000.00')).quantity == 7
    assert item.batches.get(raw_unit_cost=Decimal('2600.00')).quantity == 3


@pytest.mark.django_db
def test_general_batch_quantity_can_be_edited_but_not_below_sold(api_client, admin_user):
    item = PartInventory.objects.create(
        part_name='General starter', part_number='GST-1', category='Electrical', quantity=0, unit='piece',
    )
    created = api_client.post(
        '/api/operations/inventory-batches/',
        {'item': str(item.id), 'quantity': 5, 'raw_unit_cost': '7000.00', 'source_label': 'General inventory'},
        format='json',
    )
    batch = InventoryBatch.objects.get(id=created.data['id'])
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': batch, 'quantity': 2, 'sold_price': Decimal('10000.00')}],
    )

    valid = api_client.patch(
        f'/api/operations/inventory-batches/{batch.id}/',
        {'quantity': 4, 'raw_unit_cost': '7500.00'},
        format='json',
    )
    invalid = api_client.patch(
        f'/api/operations/inventory-batches/{batch.id}/',
        {'quantity': 1},
        format='json',
    )

    assert valid.status_code == 200
    assert valid.data['available_quantity'] == 2
    assert valid.data['net_unit_cost'] == '7500.00'
    assert invalid.status_code == 400
    batch.refresh_from_db()
    item.refresh_from_db()
    assert batch.quantity == 4
    assert item.quantity == 4


@pytest.mark.django_db
def test_general_batch_with_sales_cannot_be_deleted(api_client, admin_user):
    item = PartInventory.objects.create(part_name='General sold part', quantity=1, unit='piece')
    batch = InventoryBatch.objects.create(
        item=item, quantity=1, raw_unit_cost=Decimal('100.00'), net_unit_cost=Decimal('100.00'),
        source_label='General inventory',
    )
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': batch, 'quantity': 1, 'sold_price': Decimal('200.00')}],
    )

    response = api_client.delete(f'/api/operations/inventory-batches/{batch.id}/')

    assert response.status_code == 400
    assert InventoryBatch.objects.filter(id=batch.id).exists()


@pytest.mark.django_db
def test_container_inventory_edit_and_delete_keep_aggregate_stock_synchronized(api_client):
    container = Container.objects.create(reference='SYNC-DELETE-001')
    created = api_client.post(
        '/api/operations/items/',
        {
            'container': str(container.id),
            'part_name': 'Synchronized engine',
            'part_number': 'SYNC-A',
            'category': 'Engine',
            'quantity': 2,
            'unit': 'piece',
            'raw_unit_cost': '5000.00',
        },
        format='json',
    )
    assert created.status_code == 201
    item = ContainerItem.objects.get(id=created.data['id'])
    aggregate = PartInventory.objects.get(part_name='Synchronized engine')
    assert aggregate.quantity == 2

    updated = api_client.patch(
        f'/api/operations/items/{item.id}/',
        {'quantity': 5, 'raw_unit_cost': '5500.00'},
        format='json',
    )
    assert updated.status_code == 200
    aggregate.refresh_from_db()
    item.refresh_from_db()
    assert aggregate.quantity == 5
    assert item.inventory_batch.quantity == 5

    deleted = api_client.delete(f'/api/operations/items/{item.id}/')
    assert deleted.status_code == 204
    aggregate.refresh_from_db()
    assert aggregate.quantity == 0
    assert not ContainerItem.objects.filter(id=item.id).exists()
    assert not InventoryBatch.objects.filter(container_item_id=item.id).exists()


@pytest.mark.django_db
def test_sold_container_inventory_delete_is_blocked_without_partial_stock_change(api_client, admin_user):
    container = Container.objects.create(reference='SYNC-SOLD-001')
    created = api_client.post(
        '/api/operations/items/',
        {
            'container': str(container.id),
            'part_name': 'Protected sold engine',
            'part_number': 'PROTECTED-A',
            'category': 'Engine',
            'quantity': 3,
            'unit': 'piece',
            'raw_unit_cost': '5000.00',
        },
        format='json',
    )
    item = ContainerItem.objects.get(id=created.data['id'])
    batch = item.inventory_batch
    aggregate = batch.item
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': batch, 'quantity': 2, 'sold_price': Decimal('9000.00')}],
    )

    too_small = api_client.patch(
        f'/api/operations/items/{item.id}/', {'quantity': 1}, format='json',
    )
    changed_identity = api_client.patch(
        f'/api/operations/items/{item.id}/', {'part_number': 'CHANGED'}, format='json',
    )
    response = api_client.delete(f'/api/operations/items/{item.id}/')

    assert too_small.status_code == 400
    assert changed_identity.status_code == 400
    assert response.status_code == 409
    item.refresh_from_db()
    batch.refresh_from_db()
    aggregate.refresh_from_db()
    assert item.quantity == 3
    assert batch.quantity == 3
    assert aggregate.quantity == 3


@pytest.mark.django_db
def test_subpart_split_over_available_quantity_returns_validation_error(api_client):
    container = Container.objects.create(reference='CNT-API-SUB', status=Container.Status.READY_FOR_AUCTION)
    parent_part = PartInventory.objects.create(
        part_name='Damaged lamp',
        part_number='D-LAMP',
        category='Lights',
        quantity=1,
        unit='piece',
    )
    parent = ContainerItem.objects.create(
        container=container,
        part_name='Damaged lamp',
        part_number='D-LAMP',
        category='Lights',
        quantity=1,
        unit='piece',
        raw_unit_cost=Decimal('5000.00'),
        net_unit_cost=Decimal('5000.00'),
    )
    InventoryBatch.objects.create(
        item=parent_part,
        container=container,
        container_item=parent,
        quantity=1,
        raw_unit_cost=parent.raw_unit_cost,
        net_unit_cost=parent.net_unit_cost,
    )

    response = api_client.post(
        f'/api/operations/items/{parent.id}/subparts/',
        {
            'split_quantity': 2,
            'subparts': [
                {
                    'part_name': 'Lamp clip',
                    'part_number': 'CLIP',
                    'category': 'Lights',
                    'quantity': 1,
                    'unit': 'piece',
                    'raw_unit_cost': '100.00',
                },
            ],
        },
        format='json',
    )

    assert response.status_code == 400
    assert 'split_quantity' in response.data
    parent.refresh_from_db()
    assert parent.quantity == 1


@pytest.mark.django_db
def test_inventory_report_spreadsheet_excludes_zero_available_batches(api_client, admin_user):
    container = Container.objects.create(reference='REPORT-CNT', created_by=admin_user, updated_by=admin_user)
    item = PartInventory.objects.create(part_name='Report light', category='Lights', quantity=2, unit='piece')
    available_batch = InventoryBatch.objects.create(
        item=item,
        container=container,
        quantity=2,
        raw_unit_cost=Decimal('1000.00'),
        net_unit_cost=Decimal('1200.00'),
    )
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': available_batch, 'quantity': 1, 'sold_price': Decimal('1500.00')}],
    )
    sold_out_item = PartInventory.objects.create(part_name='Sold out report part', quantity=1, unit='piece')
    sold_out_batch = InventoryBatch.objects.create(
        item=sold_out_item,
        container=container,
        quantity=1,
        raw_unit_cost=Decimal('100.00'),
        net_unit_cost=Decimal('100.00'),
    )
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': sold_out_batch, 'quantity': 1, 'sold_price': Decimal('150.00')}],
    )

    response = api_client.get(f'/api/operations/inventory-report/?export=csv&container={container.id}')

    assert response.status_code == 200
    body = response.content.decode()
    assert 'Report light' in body
    assert 'Sold out report part' not in body
    assert 'Available quantity' in body


@pytest.mark.django_db
def test_container_tracking_fields_and_reports(api_client, admin_user):
    response = api_client.post(
        '/api/operations/containers/',
        {
            'reference': 'TRACK-CNT-001',
            'supplier_name': 'Test Agent',
            'size_type': '45 ft Custom',
            'current_location': 'Karachi Port',
            'notes': 'Tracking notes that must remain visible.',
            'status': Container.Status.READY_FOR_AUCTION,
            'added_cost': '5000.00',
        },
        format='json',
    )
    assert response.status_code == 201
    assert response.data['notes'] == 'Tracking notes that must remain visible.'

    pdf = api_client.get('/api/operations/container-tracking-report/?export=pdf')
    spreadsheet = api_client.get('/api/operations/container-tracking-report/?export=xlsx')
    assert pdf.status_code == 200
    assert pdf.content.startswith(b'%PDF')
    assert spreadsheet.status_code == 200

    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(spreadsheet.content))
    sheet = workbook['Container Tracking']
    assert [sheet.cell(row=4, column=index).value for index in range(1, 7)] == [
        'Serial No.', 'Container Number', 'Size / Type', 'Current Location', 'Agent', 'Notes',
    ]
    assert sheet['B5'].value == 'TRACK-CNT-001'


@pytest.mark.django_db
def test_auction_inventory_sheet_requires_ready_container(api_client, admin_user):
    ready = Container.objects.create(reference='AUCTION-CNT', status=Container.Status.READY_FOR_AUCTION, created_by=admin_user, updated_by=admin_user)
    not_ready = Container.objects.create(reference='LOADING-CNT', status=Container.Status.GODOWN_LOADING, created_by=admin_user, updated_by=admin_user)
    item = PartInventory.objects.create(part_name='Auction engine', category='Engine', quantity=2, unit='piece')
    InventoryBatch.objects.create(item=item, container=ready, quantity=2, raw_unit_cost=Decimal('1000.00'), net_unit_cost=Decimal('1200.00'))

    response = api_client.get(f'/api/operations/inventory-report/?report=auction&export=pdf&container={ready.id}')
    rejected = api_client.get(f'/api/operations/inventory-report/?report=auction&export=pdf&container={not_ready.id}')
    assert response.status_code == 200
    assert response.content.startswith(b'%PDF')
    assert rejected.status_code == 404

    units = _auction_inventory_units(build_inventory_report(container_id=ready.id))
    assert len(units) == 2
    assert [unit['unit_number'] for unit in units] == [1, 2]
    assert all(unit['quantity'] == 2 for unit in units)
    assert all(unit['part_name'] == 'Auction engine' for unit in units)


@pytest.mark.django_db
def test_container_profit_loss_report_uses_current_batch_costs(api_client, admin_user):
    container = Container.objects.create(
        reference='PNL-CNT-001',
        status=Container.Status.READY_FOR_AUCTION,
        added_cost=Decimal('200.00'),
        created_by=admin_user,
        updated_by=admin_user,
    )
    item = PartInventory.objects.create(part_name='P&L engine', quantity=10, unit='piece')
    batch = InventoryBatch.objects.create(
        item=item,
        container=container,
        quantity=10,
        raw_unit_cost=Decimal('100.00'),
        net_unit_cost=Decimal('120.00'),
    )
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': batch, 'quantity': 3, 'sold_price': Decimal('200.00')}],
    )

    report = build_container_profit_loss_report(container=container)
    response = api_client.get(f'/api/operations/container-profit-loss-report/?container={container.id}')
    missing = api_client.get('/api/operations/container-profit-loss-report/')

    assert report.totals['revenue'] == Decimal('600.00')
    assert report.totals['sold_cost'] == Decimal('360.00')
    assert report.totals['gross_profit'] == Decimal('240.00')
    assert report.totals['remaining_inventory_value'] == Decimal('840.00')
    assert response.status_code == 200
    assert response.content.startswith(b'%PDF')
    assert missing.status_code == 400


@pytest.mark.django_db
def test_sale_invoice_pdf_endpoint_returns_pdf(api_client, admin_user):
    customer = Customer.objects.create(name='Invoice Customer', phone='+923009999991')
    item = PartInventory.objects.create(part_name='Invoice bumper', quantity=1, unit='piece')
    batch = InventoryBatch.objects.create(item=item, quantity=1, raw_unit_cost=Decimal('5000.00'), net_unit_cost=Decimal('6000.00'))
    sale = create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CREDIT,
        customer=customer,
        lines=[{'inventory_batch': batch, 'quantity': 1, 'sold_price': Decimal('9500.00')}],
    )

    response = api_client.get(f'/api/operations/auction-sales/{sale.id}/invoice/')

    assert response.status_code == 200
    assert response['Content-Type'] == 'application/pdf'
    assert response.content.startswith(b'%PDF')

    thermal_response = api_client.get(f'/api/operations/auction-sales/{sale.id}/invoice/?layout=thermal')
    assert thermal_response.status_code == 200
    assert thermal_response['Content-Type'] == 'application/pdf'
    assert thermal_response.content.startswith(b'%PDF')
    assert 'Zulfiqar Old Spare Parts - Thermal Sales Invoice' in thermal_response['Content-Disposition']
    assert '_' not in thermal_response['Content-Disposition']


@pytest.mark.django_db
def test_sales_analytics_returns_daily_monthly_yearly_totals(api_client, admin_user):
    item = PartInventory.objects.create(part_name='Analytics mirror', quantity=2, unit='piece')
    batch = InventoryBatch.objects.create(item=item, quantity=2, raw_unit_cost=Decimal('1000.00'), net_unit_cost=Decimal('1200.00'))
    create_auction_sale(
        user=admin_user,
        sale_date=timezone.localdate(),
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': batch, 'quantity': 1, 'sold_price': Decimal('2500.00')}],
    )

    response = api_client.get('/api/operations/auction-sales/analytics/')

    assert response.status_code == 200
    assert response.data['totals']['sale_count'] >= 1
    assert response.data['daily']
    assert response.data['monthly']
    assert response.data['yearly']
    assert response.data['selected']
    assert response.data['date_bounds']['oldest_sale_date']
    assert response.data['date_bounds']['latest_sale_date']


@pytest.mark.django_db
def test_sales_analytics_supports_bounded_date_range_and_granularity(api_client, admin_user):
    item = PartInventory.objects.create(part_name='Range analytics part', quantity=2, unit='piece')
    first_batch = InventoryBatch.objects.create(item=item, quantity=1, raw_unit_cost=Decimal('100.00'), net_unit_cost=Decimal('120.00'))
    second_batch = InventoryBatch.objects.create(item=item, quantity=1, raw_unit_cost=Decimal('100.00'), net_unit_cost=Decimal('120.00'))
    first_date = timezone.localdate() - timedelta(days=60)
    second_date = timezone.localdate()
    create_auction_sale(
        user=admin_user,
        sale_date=first_date,
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': first_batch, 'quantity': 1, 'sold_price': Decimal('500.00')}],
    )
    create_auction_sale(
        user=admin_user,
        sale_date=second_date,
        payment_type=AuctionSale.PaymentType.CASH,
        lines=[{'inventory_batch': second_batch, 'quantity': 1, 'sold_price': Decimal('700.00')}],
    )

    response = api_client.get(f'/api/operations/auction-sales/analytics/?start={first_date}&end={second_date}&granularity=year')

    assert response.status_code == 200
    assert response.data['date_bounds']['start'] == first_date
    assert response.data['date_bounds']['end'] == second_date
    assert response.data['date_bounds']['granularity'] == 'year'
    assert len(response.data['selected']) >= 1
